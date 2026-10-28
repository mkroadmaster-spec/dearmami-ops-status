"""주간·월간 리포트 Word(.docx) — period_report_data 결과를 그대로 문서화 (python-docx)."""
from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.shared import Pt

from . import config
from .etl.period_report import _pct, _won

LABEL = {"gross_sales": "매출", "orders": "주문수", "aov": "객단가", "visitors": "방문자", "conversion_rate": "구매전환율",
         "ad_cost": "광고비", "roas": "ROAS", "ad_cost_ratio": "광고비율", "new_orders": "신규 주문", "returning_orders": "재구매 주문",
         "sessions": "세션", "ga4_conversions": "GA4 전환"}


def _fmt(k: str, v: float | None) -> str:
    if v is None:
        return "—"
    if k in ("gross_sales", "aov", "ad_cost", "ad_revenue", "cpc", "cpa", "sales"):
        return _won(v)
    if k in ("conversion_rate", "roas", "ad_cost_ratio", "ctr", "cvr", "cart_rate", "mobile_share"):
        return f"{v:.1f}%"
    return f"{v:,.0f}"


def _table(doc, header: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    doc.add_paragraph()


def build(data: dict[str, Any]) -> bytes:
    brand = config.sources().get("brand", "")
    pw = data["prev_word"]
    doc = Document()
    st = doc.styles["Normal"]
    st.font.size = Pt(10)
    doc.add_heading(f"{brand} {data['title']}", 0)
    doc.add_paragraph(f"기간 {data['range']['from']} ~ {data['range']['to']}  (비교: {data['prev_range']['from']} ~ {data['prev_range']['to']})")
    if data.get("missing_dates"):
        doc.add_paragraph(f"⚠ 수집 누락일: {', '.join(data['missing_dates'])}")

    n = 1
    doc.add_heading(f"{n}. 한 장 요약", 1); n += 1
    doc.add_paragraph(data["headline"])
    s = data["summary"]
    _table(doc, ["지표", "이번 기간", pw, "증감"],
           [[LABEL.get(k, k), _fmt(k, v["cur"]), _fmt(k, v["prev"]), _pct(v["pct"])] for k, v in s.items() if v["cur"] is not None])
    if data.get("better") or data.get("worse"):
        doc.add_paragraph("좋아진 것: " + (", ".join(LABEL.get(k, k) for k in data["better"]) or "—"))
        doc.add_paragraph("아쉬운 것: " + (", ".join(LABEL.get(k, k) for k in data["worse"]) or "—"))

    if data["channels"]:
        doc.add_heading(f"{n}. 매체별 광고 성과", 1); n += 1
        _table(doc, ["채널", "광고비", f"{pw} 대비", "클릭", "CTR", "CPC", "전환", "ROAS"],
               [[r["label"], _won(r["ad_cost"]), _pct(r["ad_cost_pct"]), _fmt("x", r["ad_clicks"]), _fmt("ctr", r["ctr"]),
                 _won(r["cpc"]), _fmt("x", r["ad_conversions"]), _fmt("roas", r["roas"])] for r in data["channels"]])

    fnl = data["funnel"]
    if fnl["totals"]["views"]:
        doc.add_heading(f"{n}. 구매 퍼널 병목", 1); n += 1
        doc.add_paragraph(f"병목 단계: {fnl['bottleneck'] or '없음'} (단계별 전환율을 {pw}와 대조)")
        _table(doc, ["단계", "이번 기간", pw, "증감(pt)"],
               [[st_["stage"], _fmt("cart_rate", st_["rate"]), _fmt("cart_rate", st_["prev_rate"]),
                 "—" if st_["delta_pt"] is None else f"{st_['delta_pt']:+.1f}"] for st_ in fnl["stages"]])
        if fnl["low_cart_products"]:
            doc.add_heading(f"{n}. 담기율 낮은 상품", 1); n += 1
            doc.add_paragraph("담기 0건은 '인기 없음'이 아니라 '담을 수 없음'(품절·판매중지·옵션 미설정) 신호로 해석")
            _table(doc, ["상품", "조회", "담기", "담기율", "상태"],
                   [[r["product"], _fmt("x", r["views"]), _fmt("x", r["cart_adds"]), _fmt("cart_rate", r["cart_rate"]),
                     "담을 수 없는 상태?" if r["cannot_add"] else "병목"] for r in fnl["low_cart_products"]])

    if data["best"]:
        doc.add_heading(f"{n}. 베스트 상품", 1); n += 1
        _table(doc, ["순위", "상품", "매출", f"{pw} 순위"],
               [[r["rank"], r["product"], _won(r["sales"]), r["prev_rank"] or "신규"] for r in data["best"]])

    c = data["customers"]
    doc.add_heading(f"{n}. 고객 구성", 1); n += 1
    _table(doc, ["구분", "이번 기간", pw],
           [["신규 주문", _fmt("x", c["new_orders"]), _fmt("x", c["prev_new_orders"])],
            ["재구매 주문", _fmt("x", c["returning_orders"]), _fmt("x", c["prev_returning_orders"])],
            ["모바일 매출 비중", _fmt("mobile_share", c["mobile_share"]), _fmt("mobile_share", c["prev_mobile_share"])]]
           + [[f"{m['member_type']} 주문", _fmt("x", m["orders"]), "—"] for m in c.get("members", [])])

    if data["keywords"]:
        doc.add_heading(f"{n}. 자연 검색 유입 검색어", 1); n += 1
        _table(doc, ["순위", "검색어", "유입", f"{pw} 순위"],
               [[r["rank"], r["keyword"], _fmt("x", r["visits"]), r["prev_rank"] or "신규"] for r in data["keywords"]])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
